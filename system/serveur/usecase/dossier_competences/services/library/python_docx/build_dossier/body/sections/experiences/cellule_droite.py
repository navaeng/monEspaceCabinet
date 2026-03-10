from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor, Pt

POLICE = "Arial"

def cellule_droite(table, exp):
    cell_droite = table.cell(0, 2)
    p_droite = cell_droite.paragraphs[0]


    p_droite.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run3 = p_droite.add_run(exp.get('Dates_Période', ''))

    run3.font.name = POLICE
    run3.font.size = Pt(11)
    run3.font.color.rgb = RGBColor(0x00, 0x20, 0x60)