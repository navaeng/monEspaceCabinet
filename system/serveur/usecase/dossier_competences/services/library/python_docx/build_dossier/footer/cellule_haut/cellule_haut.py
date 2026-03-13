from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


# def cellule_haut(tab):
#     p1 = tab.cell(0, 0).paragraphs[0]
#     p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
#     p1._p.append(parse_xml(
#         f'<w:sdt {nsdecls("w")}><w:sdtPr><w:shd w:fill="002060"/></w:sdtPr><w:sdtContent><w:p><w:pPr><w:jc '
#         f'w:val="center"/></w:pPr><w:r><w:rPr><w:color w:val="FFFFFF"/></w:rPr><w:fldChar '
#         f'w:fldCharType="begin"/><w:instrText>PAGE</w:instrText><w:fldChar '
#         f'w:fldCharType="end"/></w:r></w:p></w:sdtContent></w:sdt>'))

def cellule_haut(tab):
        p1 = tab.cell(0, 0).paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1._p.append(parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr><w:shd w:fill="002060"/><w:color w:val="FFFFFF"/></w:rPr>'
            f'<w:fldChar w:fldCharType="begin"/><w:instrText>PAGE</w:instrText><w:fldChar w:fldCharType="end"/></w:r>'
        ))