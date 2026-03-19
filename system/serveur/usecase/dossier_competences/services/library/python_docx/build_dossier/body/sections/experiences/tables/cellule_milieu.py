from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor, Pt

POLICE = "Calibri (corps)"

def cellule_milieu(table, exp):
    cell_milieu = table.cell(0, 2)
    p_milieu = cell_milieu.paragraphs[0]
    p_milieu.paragraph_format.space_after = Pt(0)

    p_milieu.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run3 = p_milieu.add_run(exp.get('Dates_Période', ''))

    run3.font.name = POLICE
    run3.font.size = Pt(11)
    run3.font.color.rgb = RGBColor(0x00, 0x20, 0x60)