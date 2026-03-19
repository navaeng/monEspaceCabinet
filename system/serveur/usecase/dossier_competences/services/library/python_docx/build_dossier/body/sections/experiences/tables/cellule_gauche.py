from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

COULEUR_PRINCIPALE = RGBColor(0x00, 0x20, 0x60)
POLICE = "Calibri (corps)"

def cellule_gauche(table, exp):
    cell_gauche = table.cell(0, 0)
    p_gauche = cell_gauche.paragraphs[0]
    p_gauche.paragraph_format.space_after = Pt(0)
    p_gauche.alignement = WD_ALIGN_PARAGRAPH.LEFT
    run = p_gauche.add_run(exp.get('Nom_Entreprise', ''))
    run.font.name = POLICE
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = COULEUR_PRINCIPALE

