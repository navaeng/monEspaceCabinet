from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor, Pt

POLICE = "Arial"

def cellule_milieu(table, exp):
    cell_milieu = table.cell(0, 1)
    p = cell_milieu.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    duree = exp.get('Durée_Mission', '')
    run_duree = p.add_run(duree)
    run_duree.font.name = POLICE
    run_duree.font.size = Pt(11)
    run_duree.font.color.rgb = RGBColor(0x00, 0x20, 0x60)